"""
Document processors for various file formats.
"""

import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional, Iterator
import hashlib

import structlog

logger = structlog.get_logger()


@dataclass
class ExtractedPage:
    """Single page/section of extracted content."""
    text: str
    page_number: int = 0
    confidence: float = 1.0
    metadata: dict = field(default_factory=dict)


@dataclass
class ExtractedDocument:
    """Complete extracted document."""
    pages: list[ExtractedPage]
    source_path: str
    file_hash: str
    format: str
    metadata: dict = field(default_factory=dict)

    @property
    def full_text(self) -> str:
        """Concatenate all pages."""
        return "\n\n".join(p.text for p in self.pages if p.text.strip())

    @property
    def total_chars(self) -> int:
        return sum(len(p.text) for p in self.pages)


class DocumentProcessor(ABC):
    """Base class for document processors."""

    SUPPORTED_EXTENSIONS: list[str] = []

    @abstractmethod
    def extract(self, file_path: Path) -> ExtractedDocument:
        """Extract text content from document."""
        pass

    @classmethod
    def supports(cls, file_path: Path) -> bool:
        """Check if this processor supports the file."""
        return file_path.suffix.lower() in cls.SUPPORTED_EXTENSIONS

    @staticmethod
    def compute_hash(file_path: Path) -> str:
        """Compute file hash for deduplication."""
        hasher = hashlib.sha256()
        with open(file_path, 'rb') as f:
            for chunk in iter(lambda: f.read(8192), b''):
                hasher.update(chunk)
        return hasher.hexdigest()


class TextProcessor(DocumentProcessor):
    """Process plain text and markdown files."""

    SUPPORTED_EXTENSIONS = ['.txt', '.md', '.markdown', '.text', '.rst']

    def extract(self, file_path: Path) -> ExtractedDocument:
        logger.info("processing_text", path=str(file_path))

        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            content = f.read()

        # Normalize whitespace
        content = self._normalize(content)

        return ExtractedDocument(
            pages=[ExtractedPage(text=content, page_number=1)],
            source_path=str(file_path),
            file_hash=self.compute_hash(file_path),
            format='text',
            metadata={
                'filename': file_path.name,
                'size_bytes': file_path.stat().st_size,
            },
        )

    def _normalize(self, text: str) -> str:
        """Normalize text content."""
        # Normalize line endings
        text = text.replace('\r\n', '\n').replace('\r', '\n')
        # Remove excessive newlines
        text = re.sub(r'\n{3,}', '\n\n', text)
        # Remove trailing whitespace
        text = '\n'.join(line.rstrip() for line in text.split('\n'))
        return text.strip()


class PDFProcessor(DocumentProcessor):
    """Process PDF documents."""

    SUPPORTED_EXTENSIONS = ['.pdf']

    def __init__(self, ocr_fallback: bool = True, ocr_threshold: float = 0.3):
        self.ocr_fallback = ocr_fallback
        self.ocr_threshold = ocr_threshold

    def extract(self, file_path: Path) -> ExtractedDocument:
        logger.info("processing_pdf", path=str(file_path))

        try:
            import pdfplumber
        except ImportError:
            raise ImportError("pdfplumber required for PDF processing. Install with: pip install pdfplumber")

        pages = []
        metadata = {}

        with pdfplumber.open(file_path) as pdf:
            metadata['num_pages'] = len(pdf.pages)

            for i, page in enumerate(pdf.pages):
                text = page.extract_text() or ''

                # Check if we got meaningful text. Width or height can be
                # missing on a broken page; treat that as "do not OCR".
                width = page.width or 0
                height = page.height or 0
                area = width * height * 0.001
                text_ratio = (len(text.strip()) / area) if area > 0 else 1.0

                # If text is sparse, might be scanned - try OCR
                if self.ocr_fallback and text_ratio < self.ocr_threshold:
                    ocr_text = self._ocr_page(page)
                    if len(ocr_text) > len(text):
                        text = ocr_text
                        logger.debug("used_ocr", page=i + 1)

                pages.append(ExtractedPage(
                    text=text,
                    page_number=i + 1,
                    metadata={
                        'width': page.width,
                        'height': page.height,
                    },
                ))

        return ExtractedDocument(
            pages=pages,
            source_path=str(file_path),
            file_hash=self.compute_hash(file_path),
            format='pdf',
            metadata=metadata,
        )

    def _ocr_page(self, page) -> str:
        """Run OCR on a page image."""
        try:
            import pytesseract
            from PIL import Image

            # Convert page to image
            img = page.to_image(resolution=200)
            pil_img = img.original

            # Run OCR
            text = pytesseract.image_to_string(pil_img)
            return text
        except Exception as e:
            logger.warning("ocr_failed", error=str(e))
            return ''


class DocxProcessor(DocumentProcessor):
    """Process Microsoft Word documents."""

    SUPPORTED_EXTENSIONS = ['.docx']

    def extract(self, file_path: Path) -> ExtractedDocument:
        logger.info("processing_docx", path=str(file_path))

        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx required. Install with: pip install python-docx")

        doc = Document(file_path)

        # Extract paragraphs
        paragraphs = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        # Extract tables as text
        for table in doc.tables:
            for row in table.rows:
                row_text = ' | '.join(cell.text.strip() for cell in row.cells)
                if row_text.strip(' |'):
                    paragraphs.append(row_text)

        full_text = '\n\n'.join(paragraphs)

        return ExtractedDocument(
            pages=[ExtractedPage(text=full_text, page_number=1)],
            source_path=str(file_path),
            file_hash=self.compute_hash(file_path),
            format='docx',
            metadata={
                'filename': file_path.name,
                'paragraph_count': len(doc.paragraphs),
                'table_count': len(doc.tables),
            },
        )


class HTMLProcessor(DocumentProcessor):
    """Process HTML files."""

    SUPPORTED_EXTENSIONS = ['.html', '.htm']

    def extract(self, file_path: Path) -> ExtractedDocument:
        logger.info("processing_html", path=str(file_path))

        try:
            from bs4 import BeautifulSoup
        except ImportError:
            raise ImportError("beautifulsoup4 required. Install with: pip install beautifulsoup4")

        with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
            html = f.read()

        soup = BeautifulSoup(html, 'html.parser')

        # Remove script and style elements
        for tag in soup(['script', 'style', 'nav', 'footer', 'header']):
            tag.decompose()

        # Get text
        text = soup.get_text(separator='\n')

        # Clean up whitespace
        lines = [line.strip() for line in text.splitlines()]
        text = '\n'.join(line for line in lines if line)

        return ExtractedDocument(
            pages=[ExtractedPage(text=text, page_number=1)],
            source_path=str(file_path),
            file_hash=self.compute_hash(file_path),
            format='html',
            metadata={
                'title': soup.title.get_text(strip=True) if soup.title else None,
            },
        )


class CSVProcessor(DocumentProcessor):
    """Process CSV files as one row of text per line."""

    SUPPORTED_EXTENSIONS = ['.csv']

    def extract(self, file_path: Path) -> ExtractedDocument:
        import csv

        logger.info("processing_csv", path=str(file_path))
        with open(file_path, 'r', encoding='utf-8', errors='replace', newline='') as handle:
            rows = list(csv.reader(handle))

        lines = []
        for row in rows:
            line = ' | '.join(cell.strip() for cell in row)
            if line.strip(' |'):
                lines.append(line)

        return ExtractedDocument(
            pages=[ExtractedPage(text='\n'.join(lines), page_number=1)],
            source_path=str(file_path),
            file_hash=self.compute_hash(file_path),
            format='csv',
            metadata={
                'filename': file_path.name,
                'row_count': len(rows),
            },
        )


class JSONProcessor(DocumentProcessor):
    """Process JSON documents and arrays of records."""

    SUPPORTED_EXTENSIONS = ['.json']

    def extract(self, file_path: Path) -> ExtractedDocument:
        import json

        logger.info("processing_json", path=str(file_path))
        with open(file_path, 'r', encoding='utf-8', errors='replace') as handle:
            payload = json.load(handle)

        pages = []
        if isinstance(payload, list):
            for index, item in enumerate(payload, start=1):
                pages.append(ExtractedPage(text=_json_item_text(item), page_number=index))
        elif isinstance(payload, dict) and isinstance(payload.get('text'), str):
            pages.append(ExtractedPage(text=payload['text'], page_number=1))
        else:
            pages.append(ExtractedPage(text=json.dumps(payload, ensure_ascii=False, indent=2), page_number=1))

        return ExtractedDocument(
            pages=pages or [ExtractedPage(text='', page_number=1)],
            source_path=str(file_path),
            file_hash=self.compute_hash(file_path),
            format='json',
            metadata={
                'filename': file_path.name,
                'record_count': len(pages),
            },
        )


def _json_item_text(item) -> str:
    import json

    if isinstance(item, str):
        return item
    if isinstance(item, dict):
        if isinstance(item.get('text'), str):
            return item['text']
        parts = []
        for key, value in item.items():
            if isinstance(value, (dict, list)):
                rendered = json.dumps(value, ensure_ascii=False)
            else:
                rendered = '' if value is None else str(value)
            parts.append(f"{key}: {rendered}")
        return '\n'.join(parts)
    return json.dumps(item, ensure_ascii=False)


# Registry of all processors
PROCESSORS: list[type[DocumentProcessor]] = [
    TextProcessor,
    PDFProcessor,
    DocxProcessor,
    HTMLProcessor,
    CSVProcessor,
    JSONProcessor,
]


def get_processor(file_path: Path) -> Optional[DocumentProcessor]:
    """Get appropriate processor for a file."""
    for proc_cls in PROCESSORS:
        if proc_cls.supports(file_path):
            return proc_cls()
    return None


def extract_document(file_path: Path) -> ExtractedDocument:
    """Extract document using appropriate processor."""
    processor = get_processor(file_path)
    if processor is None:
        raise ValueError(f"No processor for file type: {file_path.suffix}")
    return processor.extract(file_path)
